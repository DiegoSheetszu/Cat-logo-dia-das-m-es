from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = r"C:\Users\Diego\OneDrive\Documentos\New project\output\doc\requerimento_notificacao_confrontantes_com_conjuge.docx"


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.15):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def justify(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_signature_line(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=18, after=0)
    run = p.add_run("_" * 38)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


def set_font(run, size=12, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def add_bottom_border(paragraph, color="B7AA8B", size="10"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


doc = Document()

section = doc.sections[0]
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(3)
section.right_margin = Cm(2)

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
style.font.size = Pt(12)
style.paragraph_format.line_spacing = 1.15

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_spacing(p, after=4)
r = p.add_run("REQUERIMENTO PARA NOTIFICAÇÃO DE CONFRONTANTES")
set_font(r, size=14, bold=True)
r.font.color.rgb = RGBColor(0x2F, 0x2A, 0x24)
add_bottom_border(p)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_spacing(p, after=14)
r = p.add_run("Art. 213 da Lei n. 6.015/1973")
set_font(r, size=10)
r.italic = True
r.font.color.rgb = RGBColor(0x6B, 0x62, 0x57)

p = doc.add_paragraph()
set_paragraph_spacing(p, after=12)
r = p.add_run("Ilmo.(a) Sr.(a) Oficial(a) do Cartório de Registro de Imóveis da Comarca de Brumado/BA")
set_font(r)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

body_paragraphs = [
    "Eu, LUCIDALVA FERREIRA DIAS, brasileira, casada civilmente, portadora da cédula de identidade n. 03.540.943-60 SSP/BA, inscrita no CPF n. 657.379.705-15, residente e domiciliada na Rua Donizete Muzzi Meira, n. 37, Bairro Aviação, CEP 46.117-038, Brumado/BA, cônjuge de ELIAS PEREIRA DIAS, portador da cédula de identidade n. 04.600.458-06 e inscrito no CPF n. 484.280.945-00, venho, respeitosamente, requerer a NOTIFICAÇÃO DOS CONFRONTANTES não localizados, para fins de regularização e atualização da documentação do imóvel situado na Rua Luiz da Gama, n. 367, Bairro Novo Brumado, CEP 46.117-142, Brumado/BA, nos termos do art. 213 da Lei n. 6.015/1973.",
    "Informo que foram buscadas as assinaturas dos confrontantes do imóvel, porém não foi possível localizar os seguintes:",
]

for text in body_paragraphs:
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=10)
    r = p.add_run(text)
    set_font(r)
    justify(p)

items = [
    "1. JOVELINO JOSE DA SILVA, confrontante pelos fundos, com endereço informado na Rua Deodoro de Castro Leite, n. 177, Bairro Novo Brumado, CEP 46.117-140, Brumado/BA.",
    "2. IGREJA ASSOCIAÇÃO DAS FAMÍLIAS, confrontante pelo lado direito, situada à Rua Luiz da Gama, n. 365, CEP 46.117-142, Brumado/BA.",
]

for item in items:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(0)
    set_paragraph_spacing(p, after=8)
    r = p.add_run(item)
    set_font(r)
    justify(p)

p = doc.add_paragraph()
set_paragraph_spacing(p, after=10)
r = p.add_run("Diante do exposto, requer:")
set_font(r)
r.bold = True
justify(p)

requests = [
    "1. O recebimento do presente requerimento;",
    "2. A expedição de notificação dos confrontantes acima identificados, nos endereços informados;",
    "3. Caso não sejam localizados, seja certificada a diligência negativa e adotadas as providências cabíveis no procedimento de regularização do imóvel.",
]

for req in requests:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(0)
    set_paragraph_spacing(p, after=8)
    r = p.add_run(req)
    set_font(r)
    justify(p)

tail = [
    "Declaro, sob minha responsabilidade, que as informações prestadas são verdadeiras.",
    "Termos em que,",
    "Pede deferimento.",
]

for idx, text in enumerate(tail):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=4 if idx == 0 else 0, after=0)
    r = p.add_run(text)
    set_font(r)
    justify(p)

p = doc.add_paragraph()
set_paragraph_spacing(p, before=20, after=18)
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run("Brumado/BA, ____ de __________________ de 2026.")
set_font(r)

add_signature_line(doc)

for text in [
    "LUCIDALVA FERREIRA DIAS",
    "RG n. 03.540.943-60 SSP/BA",
    "CPF n. 657.379.705-15",
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=0)
    r = p.add_run(text)
    set_font(r)

doc.save(OUTPUT)
